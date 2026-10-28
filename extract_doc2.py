from docx import Document
doc = Document(r'C:\Users\Parhavi G.V\Downloads\survey paper capstone.docx')
with open('survey_extracted.txt','w',encoding='utf-8') as f:
    for p in doc.paragraphs:
        try:
            t = p.text
            if t.strip():
                f.write(t + '\n\n')
        except:
            pass
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text for c in row.cells]
            f.write(' | '.join(cells) + '\n')
        f.write('\n')
print('done')
