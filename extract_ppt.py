from pptx import Presentation
prs = Presentation(r'C:\Users\Parhavi G.V\Downloads\INTELLIVEST-MAS_Review1 (1) (2).pptx')
for i,s in enumerate(prs.slides):
    print('===', i, '===')
    for shape in s.shapes:
        if hasattr(shape,'text'):
            t = shape.text
            if t.strip():
                print(t[:300])
    print()
